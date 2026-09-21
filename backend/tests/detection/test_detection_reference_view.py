"""人工保存行的类型与选择规则；不读取真实参考原件。"""

import json
from data_pipeline.analysis.detection.reference_view import ReferenceView, ReferenceSource


def csv_source(name, cells):
    return ReferenceSource(
        name,
        name,
        len(cells),
        (
            {
                "source_id": name,
                "row": i,
                "location": "csv",
                "rule": "reference-rows/v2",
                "csv_record_kind": "record",
                "raw_row": json.dumps(row),
            }
            for i, row in enumerate(cells)
        ),
    )


def test_saved_reference_first_last_and_key_types_with_whole_rows(tmp_path):
    source = csv_source(
        "as_info",
        [
            ["asn", "as_name", "import_as", "is_ddos_provider", "extra"],
            ["001", "first", "[2, '3']", "True", "whole-row"],
            ["1", "second", "[]", "False", "discarded"],
        ],
    )
    view = ReferenceView("refs-v1", "observations-run:7", tmp_path / "refs")
    view.add(source)
    view.add(csv_source("important_as_dict", [["aut-num", "value"], ["001", "x"]]))
    view.add(
        csv_source(
            "triplet_info",
            [
                ["first_as", "second_as", "third_as", "stability", "appear_num"],
                ["1", "2", "3", "0.1", "1"],
                ["1", "2", "3", "0.8", "5"],
            ],
        )
    )
    assert view.mapping("as_info")["1"]["as_name"] == "first"
    assert view.mapping("as_info")["1"]["extra"] == "whole-row"
    assert view.mapping("as_info")["1"]["import_as"] == [2, "3"]
    assert view.mapping("as_info")["1"]["is_ddos_provider"] is True
    assert list(view.mapping("important_as_dict")) == [1]
    # DataFrame.apply(axis=1) 在全数值行中将 ASN 临时提升为 float；照旧保留。
    assert view.mapping("triplet_info")["1.0"]["2.0"]["3.0"]["stability"] == 0.8
    assert view.mapping("triplet_info")["1.0"]["2.0"]["3.0"]["appear_num"] == 5.0
    assert view.selection("as_info", "1")["row"] == 1
    assert view.selection("triplet_info", ("1.0", "2.0", "3.0"))["row"] == 2


def test_json_duplicate_last_and_array_types_are_preserved(tmp_path):
    view = ReferenceView("r", "fixed:1", tmp_path / "refs")
    view.add(
        ReferenceSource(
            "as_rel_dict",
            "json",
            2,
            (
                {
                    "source_id": "json",
                    "row": i,
                    "rule": "reference-rows/v2",
                    "location": "1",
                    "raw_row": json.dumps({"__object_pairs__": pairs}),
                }
                for i, pairs in enumerate(
                    [
                        [["customer", [2]]],
                        [["customer", [2]], ["customer", ["3", 4]]],
                    ]
                )
            ),
        )
    )
    assert view.mapping("as_rel_dict") == {"1": {"customer": ["3", 4]}}
    assert view.selection("as_rel_dict", "1")["row"] == 1


def test_quoted_whitespace_record_is_not_rewritten_into_blank_line(tmp_path):
    view = ReferenceView("r", "fixed:1", tmp_path / "refs")
    view.add(csv_source("important_as_dict", [["aut-num"], ["   "]]))
    assert list(view.mapping("important_as_dict")) == ["   "]


def test_country_last_row_and_prefix_v4_before_v6(tmp_path):
    def excel(role, rows):
        return ReferenceSource(
            role,
            role,
            len(rows),
            (
                {
                    "source_id": role,
                    "row": i,
                    "rule": "reference-rows/v2",
                    "location": "Sheet1",
                    "raw_row": json.dumps(
                        [
                            {
                                "type": "s" if isinstance(cell, str) else "n",
                                "value": cell,
                            }
                            for cell in row
                        ]
                    ),
                }
                for i, row in enumerate(rows)
            ),
        )

    view = ReferenceView("r", "fixed:1", tmp_path / "refs")
    view.add(
        excel(
            "country",
            [["two_letter_code", "chinese_short_name"], ["ZZ", "首"], ["ZZ", "末"]],
        )
    )
    view.add(excel("important_prefix_v4", [["prefix", "name"], ["10.0.0.0/24", "v4"]]))
    view.add(excel("important_prefix_v6", [["prefix", "name"], ["10.0.0.0/24", "v6"]]))
    assert view.mapping("country")["ZZ"]["chinese_short_name"] == "末"
    assert view.mapping("important_prefix_dict")["10.0.0.0/24"]["name"] == "v4"


def test_original_xlsx_types_survive_saved_rows_and_reconstruction(tmp_path):
    import hashlib
    import pandas as pd
    from openpyxl import Workbook, load_workbook
    from data_pipeline.bgp.input.reference_reader import rows

    original = tmp_path / "original.xlsx"
    book = Workbook()
    headers = [
        "two_letter_code",
        "formula_text",
        "error_text",
        "number_text",
        "number",
        "boolean",
        "error",
        "empty",
        "empty_text",
    ]
    book.active.append(headers)
    values = [
        ("s", "ZZ"),
        ("s", "=1+1"),
        ("s", "#N/A"),
        ("s", "001"),
        ("n", 7),
        ("b", True),
        ("e", "#N/A"),
        ("n", None),
        ("s", ""),
    ]
    for column, (kind, value) in enumerate(values, 1):
        cell = book.active.cell(2, column, value)
        cell.data_type = kind
    book.save(original)
    book.close()
    sha = hashlib.sha256(original.read_bytes()).hexdigest()
    saved = list(rows(original, sha))
    view = ReferenceView("r", "fixed:1", tmp_path / "refs")
    view.add(ReferenceSource("country", sha, len(saved), iter(saved)))
    expected = (
        pd.read_excel(original, keep_default_na=False)
        .fillna("")
        .set_index("two_letter_code")
        .to_dict("index")
    )
    assert view.mapping("country") == expected
    assert view.mapping("country")["ZZ"]["formula_text"] == "=1+1"
    assert view.mapping("country")["ZZ"]["error_text"] == "#N/A"
    rebuilt = load_workbook(tmp_path / "refs/country.xlsx")
    source_book = load_workbook(original)
    assert [(c.data_type, c.value) for c in rebuilt.active[2]] == [
        (c.data_type, c.value) for c in source_book.active[2]
    ]
    assert json.loads(saved[1]["raw_row"])[-1] == {"type": "inlineStr", "value": None}
    source_book.close()
    rebuilt.close()


def test_actual_formula_and_date_saved_rows_are_rejected(tmp_path):
    import hashlib
    from datetime import datetime
    from openpyxl import Workbook
    from data_pipeline.bgp.input.reference_reader import rows
    import pytest

    for label, value in [("formula", "=1+1"), ("date", datetime(2026, 1, 1))]:
        path = tmp_path / (label + ".xlsx")
        book = Workbook()
        book.active.append(["two_letter_code", "value"])
        book.active.append(["ZZ", value])
        book.save(path)
        book.close()
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        saved = list(rows(path, sha))
        view = ReferenceView("r", "fixed:1", tmp_path / label)
        with pytest.raises(ValueError, match="类型合同"):
            view.add(ReferenceSource("country", sha, len(saved), iter(saved)))
        assert "country" not in view.sources
        with pytest.raises(ValueError, match="11类"):
            view.take_info()


def test_unsupported_or_mismatched_saved_excel_types_cannot_be_admitted(tmp_path):
    import pytest

    for index, (kind, value) in enumerate([("str", "cached"), ("n", "7"), ("b", 1)]):
        saved = [
            {
                "source_id": "fixture",
                "rule": "reference-rows/v2",
                "row": 0,
                "location": "Sheet1",
                "raw_row": json.dumps([{"type": kind, "value": value}]),
            }
        ]
        view = ReferenceView("r", "fixed:1", tmp_path / str(index))
        with pytest.raises(ValueError, match="支持合同"):
            view.add(ReferenceSource("country", "fixture", 1, saved))
        assert "country" not in view.sources
        with pytest.raises(ValueError, match="11类"):
            view.take_info()
