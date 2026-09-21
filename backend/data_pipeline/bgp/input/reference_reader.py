"""版本化参考原件与整行保留；不运行来源表达式、不猜历史有效期。"""
import ast
import csv
import hashlib
import json
from pathlib import Path
import zipfile
from openpyxl import load_workbook

RULE='reference-rows/v2'
CSV_FIELDS=('csv_record_kind','csv_physical_start','csv_physical_end',
            'raw_byte_offset','raw_byte_length','raw_record_sha256')


class CSVLines:
    """跟踪解析器实际消费的物理行；仅保留哈希及坐标，不拼接原记录。"""
    def __init__(self, stream):
        self.stream=stream
        self.offset=0
        self.line=0

    def begin(self):
        self.start=self.offset
        self.first=self.line+1
        self.digest=hashlib.sha256()
        self.kind=None

    def __iter__(self):return self

    def __next__(self):
        value=next(self.stream)
        raw=value.encode('utf-8')
        self.digest.update(raw)
        self.offset+=len(raw)
        self.line+=1
        # BOM计入原件字节引用，但不交给CSV字段解析。
        if self.line==1 and value.startswith('\ufeff'):value=value[1:]
        body=value.removesuffix('\n').removesuffix('\r')
        kind='blank_line' if body=='' else ('whitespace_line' if all(c in ' \t' for c in body) else 'record')
        self.kind=kind if self.line==self.first else 'record'
        return value

    def metadata(self):
        return dict(csv_record_kind=self.kind,csv_physical_start=self.first,
                    csv_physical_end=self.line,raw_byte_offset=self.start,
                    raw_byte_length=self.offset-self.start,
                    raw_record_sha256=self.digest.hexdigest())



def safe_list(text):
    if text == '':return [],'legacy_empty_list'
    if len(text)>1024*1024:return None,'value_too_large'
    try:
        parsed=ast.literal_eval(text)
    except (ValueError,SyntaxError,MemoryError,RecursionError):
        return None,'invalid_literal'
    return (parsed,'parsed') if isinstance(parsed,list) else (None,'not_list')


def rows(path, expected_sha):
    """每个源行/顶层键出现保留一次，去重策略留给显式兼容消费者。"""
    path=Path(path)
    before=path.stat()
    h=hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda:f.read(1024**2),b''):h.update(b)
    if h.hexdigest()!=expected_sha:raise ValueError('参考原件SHA不符')
    common={**dict.fromkeys(CSV_FIELDS),'source_id':expected_sha,'effective_time_state':'unknown','rule':RULE}
    if zipfile.is_zipfile(path):
        # 按magic识别OOXML，兼容错误.xls后缀；原单元格类型经JSON标记保留。
        with path.open('rb') as f:
            book=load_workbook(f,read_only=True,data_only=False)
            try:
                for sheet in book:
                    for i,row in enumerate(sheet.iter_rows()):
                        values=[{'type':c.data_type,'value':c.value} for c in row]
                        yield {**common,'row':i,'location':sheet.title,'raw_row':json.dumps(values,ensure_ascii=False,default=str)}
            finally:book.close()
    elif path.suffix.lower()=='.csv':
        # UTF-8字段字符数不可能超过已固定原件的总字节数；进程另受RSS约束。
        previous=csv.field_size_limit(max(1,before.st_size))
        try:
            with path.open(encoding='utf-8',newline='') as f:
                lines=CSVLines(f)
                reader=csv.reader(lines,strict=True)
                i=0
                while True:
                    lines.begin()
                    try:row=next(reader)
                    except StopIteration:break
                    except (csv.Error,UnicodeError) as exc:
                        raise ValueError(f'CSV参考解析失败 source={expected_sha} path={path} row={i} physical_start={lines.first} physical_end={lines.line} byte_offset={lines.start}: {exc}') from exc
                    yield {**common,**lines.metadata(),'row':i,'location':'csv','raw_row':json.dumps(row,ensure_ascii=False)}
                    i+=1
        finally:csv.field_size_limit(previous)
    else:
        if before.st_size>128*1024**2:raise ValueError('JSON参考超过显式内存上限')
        # object_pairs_hook保留重复键及顺序，避免默认dict静默覆盖。
        with path.open(encoding='utf-8') as f:
            value=json.load(f,object_pairs_hook=lambda pairs:{'__object_pairs__':pairs})
        if not isinstance(value,dict) or '__object_pairs__' not in value:raise ValueError('参考顶层必须为对象')
        for i,(key,value) in enumerate(value['__object_pairs__']):
            yield {**common,'row':i,'location':key,'raw_row':json.dumps(value,ensure_ascii=False)}
    after=path.stat()
    if (before.st_size,before.st_mtime_ns,before.st_ctime_ns,before.st_ino)!=(after.st_size,after.st_mtime_ns,after.st_ctime_ns,after.st_ino):
        raise ValueError('读取期间参考原件变化')


def resolved_interpretation(path):
    """别名可以换位置，不能在同一计划改变真实解析分支。"""
    path=Path(path)
    if zipfile.is_zipfile(path):
        return dict(format='ooxml_zip',rule=RULE,options=dict(read_only=True,data_only=False,cell_type=True,order='sheet,row',serializer='json-default-str/v1'))
    if path.suffix.lower()=='.csv':
        return dict(format='csv',rule=RULE,options=dict(encoding='utf-8',newline='',delimiter=',',quotechar='"',doublequote=True,escapechar=None,skipinitialspace=False,strict=True,physical_rows='CSVLines/v2',bom='preserve'))
    return dict(format='json',rule=RULE,options=dict(encoding='utf-8',top_level='object',pairs='preserve-duplicates-order/v1'))
